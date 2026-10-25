"""Ausencia de cobrancas iFood nunca representa margem completa na DRE/PE."""

import asyncio
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest

from app.dashboard import ponto_equilibrio_margem as pe
from app.dashboard import ponto_equilibrio_routes as pe_routes
from app import dre_export_routes
from app.dre_canais import agregacao, detalhes, routes
from app.dre_canais.base import _normalizar_canal, _novo_canal, _snapshot_pronto
from app.dre_canais.linhas import montar_linhas_dre_competencia
from app.dre_canais.schemas import DREPorCanalResponse
from app.services.rentabilidade_pendencias import aplicar_pendencia_ifood


def _venda(venda_id=1, canal="ifood"):
    snapshot = {
        "snapshot_version": 6,
        "venda_id": venda_id,
        "numero_venda": f"DEV-{venda_id}",
        "venda_bruta": 100,
        "taxa_loja": 8,
        "desconto": 0,
        "imposto": 4,
        "taxa_entrega": 3,
        "taxa_operacional": 2,
        "taxa_cartao": 5,
        "comissao": 0,
        "custo_campanha": 0,
        "custo_produtos": 40,
        "venda_liquida": 94,
        "lucro": 54,
        "margem_sobre_venda": 54,
        "itens": [],
    }
    return SimpleNamespace(
        id=venda_id,
        numero_venda=f"DEV-{venda_id}",
        data_venda=datetime(2026, 10, 24, 12),
        canal=canal,
        status="finalizada",
        desconto_venda_valor=0,
        desconto_valor=0,
        cupom_code=None,
        total=108,
        itens=[],
        pagamentos=[],
        cliente=None,
        rentabilidade_snapshot=snapshot,
    )


def _preloads(monkeypatch, module):
    monkeypatch.setattr(module, "_formas_pagamento_map", lambda *args: {})
    monkeypatch.setattr(module, "_impostos_percentual", lambda *args: 0)
    for name in (
        "_bulk_comissoes_por_venda",
        "_bulk_cupons_por_venda",
        "_bulk_cashback_por_venda",
        "_bulk_taxa_operacional_por_venda",
        "_bulk_estoque_custos_por_venda",
    ):
        monkeypatch.setattr(module, name, lambda *args: {})


def test_snapshot_dre_legado_e_protegido_sem_mutar_fotografia_congelada():
    venda = _venda()
    original = deepcopy(venda.rentabilidade_snapshot)
    protegido = _snapshot_pronto(venda)

    assert protegido["lucro"] is None
    assert protegido["taxa_cartao"] is None
    assert protegido["taxas_ifood_pendentes"] is True
    assert protegido["custo_produtos"] == 40
    assert venda.rentabilidade_snapshot == original
    assert _normalizar_canal("iFood") == "ifood"


def test_dre_mista_preserva_receita_cmv_e_nao_publica_lucro_final():
    loja = _novo_canal()
    loja.update(
        receita_produtos=Decimal("50"), cmv=Decimal("20"), taxas_cartao=Decimal("2")
    )
    ifood = _novo_canal()
    ifood.update(
        receita_produtos=Decimal("100"),
        cmv=Decimal("40"),
        vendas_rentabilidade_pendente=1,
        pendencias_rentabilidade=["Comissao e taxa do iFood aguardam conciliacao."],
    )
    linhas, totais = montar_linhas_dre_competencia(
        {"loja_fisica": loja, "ifood": ifood}
    )

    assert totais["receita_bruta"] == 150
    assert totais["cmv"] == 60
    assert totais["lucro_liquido"] is None
    assert totais["margem_liquida"] is None
    assert totais["despesas_variaveis"] is None
    assert totais["despesas_variaveis_conhecidas"] == 2
    assert totais["vendas_rentabilidade_pendente"] == 1
    assert all(linha.valor is None for linha in linhas if linha.tipo == "lucro")
    taxas = {linha.canal: linha for linha in linhas if linha.campo == "taxas_cartao"}
    assert taxas["ifood"].valor is None
    assert taxas["loja_fisica"].valor == 2


def test_agregacao_e_resposta_dre_propagam_pendencia_e_canal_correto(monkeypatch):
    venda = _venda()
    db = MagicMock()
    db.query.return_value.options.return_value.filter.return_value.all.return_value = [
        venda
    ]
    _preloads(monkeypatch, agregacao)

    dados = agregacao.obter_vendas_por_canal(db, 10, 2026, "tenant-dev")
    assert set(dados) == {"ifood"}
    assert dados["ifood"]["vendas_rentabilidade_pendente"] == 1
    assert dados["ifood"]["receita_produtos"] == 100
    assert dados["ifood"]["cmv"] == 40

    monkeypatch.setattr(routes, "obter_vendas_por_canal", lambda *args, **kwargs: dados)
    for name in (
        "agregar_devolucoes_por_canal",
        "agregar_contas_receber_manuais_por_canal",
        "agregar_contas_pagar_por_canal",
        "agregar_fretes_sobre_compras",
    ):
        monkeypatch.setattr(routes, name, lambda *args, **kwargs: None)
    resposta = routes.gerar_dre_por_canais(
        ano=2026,
        mes=10,
        mes_inicial=10,
        data_final=None,
        canais="ifood",
        db=db,
        user_and_tenant=(None, "tenant-dev"),
    )
    assert resposta.taxas_ifood_pendentes is True
    assert resposta.totais["lucro_liquido"] is None
    assert resposta.alertas[-1].codigo == "ifood_financial_pendente"


@pytest.mark.parametrize("campo", ["taxas_cartao", "taxas_marketplace"])
def test_detalhe_taxa_pendente_mantem_item_e_total_nulos(monkeypatch, campo):
    venda = _venda()
    db = MagicMock()
    db.query.return_value.options.return_value.filter.return_value.all.return_value = [
        venda
    ]
    monkeypatch.setattr(
        detalhes,
        "_preparar_snapshots_vendas",
        lambda *args: {1: aplicar_pendencia_ifood(venda, venda.rentabilidade_snapshot)},
    )

    items = detalhes._detalhes_vendas_campo(db, 10, 2026, "tenant-dev", "ifood", campo)
    assert len(items) == 1
    assert items[0].valor is None
    assert items[0].meta["taxas_ifood_pendentes"] is True
    monkeypatch.setattr(
        detalhes, "_detalhes_vendas_campo", lambda *args, **kwargs: items
    )
    monkeypatch.setattr(detalhes, "_detalhes_contas_campo", lambda *args, **kwargs: [])
    resposta = detalhes.detalhar_linha_dre_por_canal(
        campo=campo,
        canal="ifood",
        ano=2026,
        mes=10,
        mes_inicial=10,
        data_final=None,
        page=1,
        page_size=30,
        db=db,
        user_and_tenant=(None, "tenant-dev"),
    )
    assert resposta.total is None
    assert resposta.taxas_ifood_pendentes is True


def test_pe_misto_preserva_faturamento_e_mascara_taxas_margem_e_subtotal():
    venda = _venda()
    protegido = aplicar_pendencia_ifood(venda, venda.rentabilidade_snapshot)
    resultado = pe._somar_componentes_margem_vendas_pe(
        [
            {"snapshot": protegido},
            {"snapshot": _venda(2, "loja_fisica").rentabilidade_snapshot},
        ]
    )
    assert resultado["faturamento"] == 216
    assert resultado["cmv_estimado"] == 80
    assert resultado["taxas_cartao"] is None
    assert resultado["taxas_cartao_conhecidos"] == 5
    assert resultado["margem_contribuicao"] is None
    assert resultado["margem_decimal"] is None
    taxas = next(
        item
        for item in resultado["detalhes_margem"]["subtotais"]
        if item["id"] == "taxas_cartao"
    )
    assert taxas["valor"] is None
    assert (
        resultado["detalhes_margem"]["componentes"]["taxas_cartao"][0]["valor"] is None
    )


def test_modo_fiscal_nao_recria_lucro_a_partir_de_snapshot_pendente():
    venda = _venda()
    protegido = aplicar_pendencia_ifood(venda, venda.rentabilidade_snapshot)
    ajustado = pe._ajustar_snapshot_custo_fiscal_pe(
        protegido, venda_tem_documento=False, modo_custo_fiscal="documentos_emitidos"
    )
    assert ajustado["imposto"] == 0
    assert ajustado["lucro"] is None
    assert ajustado["venda_liquida"] is None


def test_detalhes_pe_nao_convertem_total_ou_item_pendente_em_zero(monkeypatch):
    venda = _venda()
    resultado = pe._somar_componentes_margem_vendas_pe(
        [{"snapshot": aplicar_pendencia_ifood(venda, venda.rentabilidade_snapshot)}]
    )
    resultado["periodo"] = {"inicio": date(2026, 10, 1), "fim": date(2026, 10, 31)}

    async def resposta_pe(**kwargs):
        return resultado

    monkeypatch.setattr(pe_routes, "obter_ponto_equilibrio", resposta_pe)
    resposta = asyncio.run(
        pe_routes.obter_ponto_equilibrio_detalhes(
            grupo="taxas_cartao",
            db=object(),
            user_and_tenant=(None, "tenant-dev"),
        )
    )
    assert resposta["total"] is None
    assert resposta["items"][0]["valor"] is None
    assert resposta["taxas_ifood_pendentes"] is True


def test_exports_pdf_excel_preservam_pendente_sem_zero_nem_erro(monkeypatch):
    import pdfplumber
    from openpyxl import load_workbook

    canal = _novo_canal()
    canal.update(
        receita_produtos=Decimal("100"),
        cmv=Decimal("40"),
        vendas_rentabilidade_pendente=1,
        pendencias_rentabilidade=["Aguardando Financial."],
    )
    linhas, totais = montar_linhas_dre_competencia({"ifood": canal})
    dre = DREPorCanalResponse(
        periodo="Outubro 2026",
        mes_inicial=10,
        mes=10,
        ano=2026,
        linhas=linhas,
        totais=totais,
        canais_encontrados=["ifood"],
        rentabilidade_status="pending_ifood_financial",
        taxas_ifood_pendentes=True,
        vendas_rentabilidade_pendente=1,
        pendencias_rentabilidade=["Aguardando Financial."],
    )
    monkeypatch.setattr(dre_export_routes, "gerar_dre_por_canais", lambda **kwargs: dre)
    parametros = dict(
        ano=2026,
        mes=10,
        mes_inicial=10,
        data_final=None,
        canais="ifood",
        db=object(),
        user_and_tenant=(None, "tenant-dev"),
    )
    dto = dre_export_routes._dre_para_exportacao(**parametros)
    assert dto.taxas_cartao is None
    assert dto.lucro_liquido is None
    assert dto.margem_liquida is None
    assert dto.receita_bruta == Decimal("100")

    async def conteudo(resposta):
        return b"".join([parte async for parte in resposta.body_iterator])

    excel = asyncio.run(dre_export_routes.exportar_dre_excel(**parametros))
    planilha = load_workbook(
        BytesIO(asyncio.run(conteudo(excel))), read_only=True, data_only=True
    )
    try:
        valores = {
            linha[0]: linha[1]
            for linha in planilha["DRE"].iter_rows(values_only=True)
            if linha[0]
        }
        assert valores["RECEITA BRUTA"] == 100
        assert valores["(-) CMV (Custo Mercadorias Vendidas)"] == 40
        assert valores["  Taxas de Cartão"] == "Pendente iFood"
        assert valores["(=) LUCRO LÍQUIDO"] == "Pendente iFood"
    finally:
        planilha.close()

    pdf = asyncio.run(dre_export_routes.exportar_dre_pdf(**parametros))
    with pdfplumber.open(BytesIO(asyncio.run(conteudo(pdf)))) as documento:
        texto = "\n".join(pagina.extract_text() or "" for pagina in documento.pages)
    assert "Pendente iFood" in texto
    assert "R$ 100,00" in texto
    assert "LUCRO LÍQUIDO Pendente iFood" in texto


@pytest.mark.parametrize("referencia_pendente", [False, True])
def test_rota_pe_nao_reclassifica_pendencia_como_margem_insuficiente(
    monkeypatch, referencia_pendente
):
    venda = _venda()
    pendente = pe._somar_componentes_margem_vendas_pe(
        [{"snapshot": aplicar_pendencia_ifood(venda, venda.rentabilidade_snapshot)}]
    )
    completo = pe._somar_componentes_margem_vendas_pe(
        [{"snapshot": _venda(2, "loja_fisica").rentabilidade_snapshot}]
    )
    for resultado in (pendente, completo):
        resultado.update(quantidade_vendas=1, ticket_medio=108)
    periodo = completo if referencia_pendente else pendente
    referencia = pendente if referencia_pendente else completo
    db = MagicMock()
    db.query.return_value.join.return_value.outerjoin.return_value.filter.return_value.scalar.return_value = 0
    db.query.return_value.outerjoin.return_value.outerjoin.return_value.outerjoin.return_value.outerjoin.return_value.filter.return_value.all.return_value = []
    monkeypatch.setattr(
        pe_routes,
        "inspect",
        lambda *args: SimpleNamespace(has_table=lambda *args: False),
    )
    monkeypatch.setattr(
        pe_routes,
        "_calcular_folha_gerencial_estimada",
        lambda *args: {"total": 0, "quantidade_funcionarios": 0},
    )
    monkeypatch.setattr(
        pe_routes,
        "_calcular_margem_periodo_ponto_equilibrio",
        lambda *args, **kwargs: periodo,
    )
    monkeypatch.setattr(
        pe_routes,
        "_calcular_margem_referencia_ponto_equilibrio",
        lambda *args, **kwargs: referencia,
    )
    resultado = asyncio.run(
        pe_routes.obter_ponto_equilibrio(
            data_inicio=date(2026, 10, 1),
            data_fim=date(2026, 10, 31),
            fonte_margem="mes_anterior_fechado",
            db=db,
            user_and_tenant=(None, "tenant-dev"),
        )
    )
    assert resultado["status"] == "pending_ifood_financial"
    assert resultado["faturamento"] == 108
    assert resultado["ponto_equilibrio"] is None
    assert resultado["falta_faturar"] is None
    assert resultado["percentual_atingido"] is None
    assert resultado["vendas_necessarias"] is None
    assert resultado["vendas_rentabilidade_pendente"] == 1
